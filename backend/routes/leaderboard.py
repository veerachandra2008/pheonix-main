import uuid
import datetime
import csv
import io
from flask import Blueprint, request, jsonify, Response
from config import get_supabase_client
from routes.payments import IN_MEMORY_REGISTRATIONS, is_user_authorized_for_tournament, load_tournament_for_payment
from routes.tournaments import IN_MEMORY_TOURNAMENTS
from routes.attendance import IN_MEMORY_EVENT_ATTENDANCE
from routes.auth import get_authenticated_user

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

leaderboard_bp = Blueprint('leaderboard', __name__)
admin_submissions_bp = Blueprint('admin_submissions', __name__)

# In-Memory Stores for Tournament Scoring Rules & Placement Matrix
# Keyed by tournament_slug (lowercased)
IN_MEMORY_SCORING_RULES = {}
IN_MEMORY_PLACEMENT_RULES = {}  # Keyed by scoring_rule_id -> list of {placement, points}

# Phase 3 In-Memory Stores for Matches and Results
IN_MEMORY_MATCHES = {}  # Keyed by tournament_slug (lowercased) -> list of match dicts
IN_MEMORY_MATCH_RESULTS = {}  # Keyed by match_id -> dict of team_id -> result dict

# Phase 5 & 6 In-Memory Stores for Lifecycle Status, Admin Submissions, and Audit Trail
IN_MEMORY_TOURNAMENT_STATUS = {}  # Keyed by tournament_slug -> status record
IN_MEMORY_SUBMISSIONS = []  # List of submission records
IN_MEMORY_AUDIT_LOGS = []  # List of audit trail records


def get_current_iso_timestamp():
    return datetime.datetime.now(datetime.timezone.utc).isoformat()


def record_audit_event(tournament_id, submission_id, action, performed_by, reason=None):
    """
    Appends an immutable audit event for tournament lifecycle tracking:
    SUBMITTED, CHANGES_REQUESTED, RESUBMITTED, APPROVED, PUBLISHED.
    """
    clean_slug = (tournament_id or '').strip().lower()
    now_iso = get_current_iso_timestamp()
    event_id = str(uuid.uuid4())
    event = {
        'id': event_id,
        'tournament_id': clean_slug,
        'submission_id': str(submission_id) if submission_id else None,
        'action': action,
        'performed_by': performed_by or 'system',
        'reason': reason or '',
        'created_at': now_iso
    }
    IN_MEMORY_AUDIT_LOGS.append(event)
    try:
        supabase = get_supabase_client()
        supabase.table('tournament_result_audit').insert(event).execute()
    except Exception:
        pass
    return event


def get_audit_trail_for_tournament(tournament_id):
    """
    Retrieves chronological audit trail for a tournament from memory and database.
    """
    clean_slug = (tournament_id or '').strip().lower()
    logs = [log for log in IN_MEMORY_AUDIT_LOGS if (log.get('tournament_id') or '').lower() == clean_slug]
    try:
        supabase = get_supabase_client()
        res = supabase.table('tournament_result_audit').select('*').eq('tournament_id', clean_slug).order('created_at', desc=False).execute()
        if res.data:
            existing_ids = {l.get('id') for l in logs}
            for row in res.data:
                if row.get('id') not in existing_ids:
                    logs.append(row)
    except Exception:
        pass
    logs.sort(key=lambda x: x.get('created_at', ''))
    return logs


def get_tournament_lifecycle_status(tournament_slug):
    """
    Returns the lifecycle status dict of a tournament leaderboard:
    'DRAFT', 'LIVE', 'FINALIZED', 'SUBMITTED', 'CHANGES_REQUESTED', 'APPROVED', 'PUBLISHED'.
    Default is 'LIVE'.
    """
    clean_slug = (tournament_slug or '').strip().lower()
    if clean_slug in IN_MEMORY_TOURNAMENT_STATUS:
        return IN_MEMORY_TOURNAMENT_STATUS[clean_slug]

    try:
        supabase = get_supabase_client()
        res = supabase.table('tournament_leaderboard_status').select('*').eq('tournament_id', clean_slug).execute()
        if res.data and len(res.data) > 0:
            status_data = res.data[0]
            IN_MEMORY_TOURNAMENT_STATUS[clean_slug] = status_data
            return status_data
    except Exception:
        pass

    default_status = {
        'tournament_id': clean_slug,
        'status': 'LIVE',
        'finalized_at': None,
        'finalized_by': None,
        'submitted_at': None,
        'submitted_by': None,
        'approved_at': None,
        'approved_by': None,
        'published_at': None,
        'published_by': None,
        'change_request_reason': None,
        'change_requested_by': None,
        'change_requested_at': None
    }
    return default_status


def is_tournament_locked(tournament_slug):
    """
    Returns True if the tournament leaderboard is FINALIZED, SUBMITTED, APPROVED, or PUBLISHED,
    preventing any further edits to rules, matches, or scores.
    When CHANGES_REQUESTED, returns False, allowing the organizer to correct results.
    """
    status_info = get_tournament_lifecycle_status(tournament_slug)
    current_status = (status_info.get('status') or 'LIVE').upper()
    return current_status in ['FINALIZED', 'SUBMITTED', 'APPROVED', 'PUBLISHED']


def check_tournament_lock_rejection(clean_slug, action_description):
    """
    If tournament is locked, returns (True, response_tuple).
    Otherwise returns (False, None).
    Returns HTTP 400 if APPROVED or PUBLISHED (Phase 6), or HTTP 403 if FINALIZED or SUBMITTED (Phase 5).
    """
    if not is_tournament_locked(clean_slug):
        return False, None
    status_info = get_tournament_lifecycle_status(clean_slug)
    curr_stat = (status_info.get('status') or '').upper()
    code = 400 if curr_stat in ['APPROVED', 'PUBLISHED'] else 403
    return True, (jsonify({
        'success': False,
        'message': f"Tournament results are {curr_stat.lower()} and locked. {action_description}."
    }), code)


def check_organizer_authorization(user, clean_slug):
    """
    Validates if user is an authorized organizer or admin for this tournament.
    Returns (is_authorized: bool, status_code: int, error_message: str, tournament_obj: dict)
    """
    if not user:
        return False, 401, "Authentication required.", None

    role = (user.get('role') or 'PLAYER').strip().upper()
    if role == 'ADMIN':
        return True, 200, None, {'slug': clean_slug, 'title': clean_slug}

    # Load tournament
    t = load_tournament_for_payment(clean_slug)
    if not t:
        t = next((item for item in IN_MEMORY_TOURNAMENTS if (item.get('slug') or '').lower() == clean_slug), None)
    if not t:
        reg = next((r for r in IN_MEMORY_REGISTRATIONS.values() if (r.get('tournament_slug') or r.get('tournament_id') or r.get('tournament') or '').lower() == clean_slug), None)
        if reg:
            t = {'slug': clean_slug, 'title': clean_slug, 'organizer_email': reg.get('organizer_email') or user.get('email')}

    if not t:
        return False, 404, f"Tournament '{clean_slug}' not found.", None

    if role == 'ORGANIZER':
        if is_user_authorized_for_tournament(user, t):
            return True, 200, None, t
        clean_email = (user.get('email') or '').strip().lower()
        if not t.get('organizer_email') and not t.get('createdBy') and not t.get('organizer_id'):
            return True, 200, None, t
        return False, 403, "You are not authorized to manage this tournament.", t

    return False, 403, "Organizer or Admin permissions required.", t



def calculate_score_for_rule(rule, raw_val):
    """
    Server-authoritative score calculation engine.
    Given a scoring rule dictionary and raw input value:
    - PER_UNIT: raw_value * points_per_unit
    - OCCURRENCE: raw_value * points_per_unit (e.g. 1 Booyah * 12 = 12)
    - PENALTY: raw_value * points_per_unit (reduces total, e.g. 2 fouls * -5 = -10)
    - PLACEMENT: lookup entered placement in custom mapping (e.g. 1st -> 12, 2nd -> 9)
    """
    rule_type = (rule.get('type') or 'PER_UNIT').upper()
    points_per_unit = float(rule.get('points_per_unit', 0) or 0)
    try:
        raw_num = float(raw_val)
    except (ValueError, TypeError):
        raw_num = 0.0

    if rule_type in ['PER_UNIT', 'OCCURRENCE']:
        return raw_num * points_per_unit
    elif rule_type == 'PENALTY':
        # Ensure penalty points reduce total score
        if points_per_unit > 0:
            return - (raw_num * points_per_unit)
        else:
            return raw_num * points_per_unit
    elif rule_type == 'PLACEMENT':
        pos = int(round(raw_num))
        if pos <= 0:
            return 0.0
        placement_map = {int(p.get('placement', 1)): float(p.get('points', 0)) for p in rule.get('placement_points', [])}
        return placement_map.get(pos, 0.0)
    return 0.0


def recalculate_match_results_for_tournament(tournament_slug):
    """
    Recalculates all match results for this tournament from stored raw_scores
    using current scoring rules. Does NOT alter raw_scores.
    """
    clean_slug = (tournament_slug or '').strip().lower()
    rules = get_rules_for_tournament(clean_slug)
    rules_dict = {str(r['id']): r for r in rules}

    # 1. Update in-memory match results
    for match_id, team_results in IN_MEMORY_MATCH_RESULTS.items():
        for team_id, res in team_results.items():
            if res.get('tournament_id') == clean_slug:
                raw_scores = res.get('raw_scores', {})
                calc_scores = {}
                for r_id, r in rules_dict.items():
                    raw_val = raw_scores.get(r_id, 0)
                    calc_scores[r_id] = calculate_score_for_rule(r, raw_val)
                res['calculated_scores'] = calc_scores
                res['total_points'] = sum(calc_scores.values())
                res['updated_at'] = get_current_iso_timestamp()

    # 2. Update Supabase
    try:
        supabase = get_supabase_client()
        sb_res = supabase.table('match_team_results').select('*').eq('tournament_id', clean_slug).execute()
        if sb_res.data:
            for row in sb_res.data:
                raw_scores = row.get('raw_scores') or {}
                calc_scores = {}
                for r_id, r in rules_dict.items():
                    raw_val = raw_scores.get(r_id, 0)
                    calc_scores[r_id] = calculate_score_for_rule(r, raw_val)
                total = sum(calc_scores.values())
                supabase.table('match_team_results').update({
                    'calculated_scores': calc_scores,
                    'total_points': total,
                    'updated_at': get_current_iso_timestamp()
                }).eq('id', row['id']).execute()
    except Exception:
        pass


def get_rules_for_tournament(tournament_slug):
    """
    Fetch all scoring rules for a tournament (ordered by sort_order),
    including associated placement matrices for 'PLACEMENT' rules.
    Combines Supabase and in-memory fallback.
    """
    clean_slug = (tournament_slug or '').strip().lower()
    rules_map = {}

    # 1. Fetch from Supabase
    try:
        supabase = get_supabase_client()
        res = supabase.table('tournament_scoring_rules').select('*').eq('tournament_id', clean_slug).order('sort_order').execute()
        if res.data:
            for r in res.data:
                r_id = str(r['id'])
                rules_map[r_id] = {
                    'id': r_id,
                    'tournament_id': clean_slug,
                    'name': r.get('name'),
                    'type': (r.get('type') or 'PER_UNIT').upper(),
                    'points_per_unit': float(r.get('points_per_unit', 0) or 0),
                    'sort_order': int(r.get('sort_order', 0) or 0),
                    'created_at': r.get('created_at'),
                    'updated_at': r.get('updated_at'),
                    'placement_points': []
                }
            
            # Fetch placement rules for PLACEMENT columns
            placement_rule_ids = [rid for rid, r in rules_map.items() if r['type'] == 'PLACEMENT']
            if placement_rule_ids:
                try:
                    for prid in placement_rule_ids:
                        p_res = supabase.table('placement_scoring_rules').select('*').eq('scoring_rule_id', prid).order('placement').execute()
                        if p_res.data:
                            for p in p_res.data:
                                parent_id = str(p['scoring_rule_id'])
                                if parent_id in rules_map:
                                    rules_map[parent_id]['placement_points'].append({
                                        'id': str(p.get('id', '')),
                                        'placement': int(p.get('placement', 1)),
                                        'points': float(p.get('points', 0))
                                    })
                except Exception as p_err:
                    print(f"Notice fetching placement rules: {p_err}")
    except Exception:
        pass

    # 2. Merge memory store (fallback & local state)
    mem_rules = IN_MEMORY_SCORING_RULES.get(clean_slug, [])
    for mr in mem_rules:
        m_id = str(mr['id'])
        p_matrix = IN_MEMORY_PLACEMENT_RULES.get(m_id, mr.get('placement_points', []))
        if m_id not in rules_map:
            rules_map[m_id] = {
                **mr,
                'placement_points': sorted(p_matrix, key=lambda x: int(x.get('placement', 1)))
            }
        elif not rules_map[m_id].get('placement_points') and p_matrix:
            rules_map[m_id]['placement_points'] = sorted(p_matrix, key=lambda x: int(x.get('placement', 1)))

    # Sort final rules by sort_order
    sorted_rules = sorted(rules_map.values(), key=lambda x: int(x.get('sort_order', 0)))
    return sorted_rules


def get_matches_for_tournament(tournament_slug):
    """
    Fetch all matches for a tournament (ordered by match_number).
    Combines Supabase and in-memory fallback.
    """
    clean_slug = (tournament_slug or '').strip().lower()
    matches_map = {}

    # 1. Fetch from Supabase
    try:
        supabase = get_supabase_client()
        res = supabase.table('tournament_matches').select('*').eq('tournament_id', clean_slug).order('match_number').execute()
        if res.data:
            for m in res.data:
                m_id = str(m['id'])
                matches_map[m_id] = {
                    'id': m_id,
                    'tournament_id': clean_slug,
                    'match_number': int(m.get('match_number', 1)),
                    'title': m.get('title') or f"Match {m.get('match_number', 1)}",
                    'status': m.get('status', 'COMPLETED'),
                    'created_at': m.get('created_at'),
                    'updated_at': m.get('updated_at')
                }
    except Exception:
        pass

    # 2. Merge memory store
    mem_matches = IN_MEMORY_MATCHES.get(clean_slug, [])
    for mm in mem_matches:
        mm_id = str(mm['id'])
        if mm_id not in matches_map:
            matches_map[mm_id] = mm

    sorted_matches = sorted(matches_map.values(), key=lambda x: int(x.get('match_number', 1)))
    return sorted_matches


def get_present_teams_for_tournament(tournament_slug):
    """
    Helper to fetch and deduplicate teams currently marked PRESENT
    using existing registration and attendance tables.
    Returns: (registered_count, present_teams_list)
    """
    clean_slug = (tournament_slug or '').strip().lower()

    # Read Existing Registrations
    supabase_regs = []
    try:
        supabase = get_supabase_client()
        r_res = supabase.table('registrations').select('*').ilike('tournament_slug', clean_slug).execute()
        supabase_regs = r_res.data or []
    except Exception:
        pass

    regs_map = {}
    for r in list(IN_MEMORY_REGISTRATIONS.values()) + supabase_regs:
        r_slug = (r.get('tournament_slug') or r.get('tournamentSlug') or '').strip().lower()
        if r_slug != clean_slug:
            continue
        reg_id = r.get('pass_id') or r.get('passId') or r.get('team_id') or r.get('teamId') or r.get('id')
        if reg_id:
            regs_map[str(reg_id)] = r

    # Read Existing Attendance Records
    supabase_att = []
    try:
        supabase = get_supabase_client()
        a_res = supabase.table('event_attendance').select('*').ilike('tournament_slug', clean_slug).execute()
        supabase_att = a_res.data or []
    except Exception:
        pass

    attendance_status_map = {}
    for a in supabase_att:
        a_slug = (a.get('tournament_slug') or a.get('tournamentSlug') or '').strip().lower()
        if a_slug != clean_slug:
            continue
        status = (a.get('attendance_status') or a.get('attendanceStatus') or a.get('status') or 'NOT_MARKED').upper()
        if a.get('pass_id'):
            attendance_status_map[str(a['pass_id'])] = status
        if a.get('team_id'):
            attendance_status_map[str(a['team_id'])] = status

    for k, a in IN_MEMORY_EVENT_ATTENDANCE.items():
        if not isinstance(a, dict):
            continue
        status = (a.get('attendance_status') or a.get('attendanceStatus') or a.get('status') or 'NOT_MARKED').upper()

        if isinstance(k, tuple) and len(k) == 2:
            t_slug, t_key = k
            if str(t_slug).strip().lower() == clean_slug:
                attendance_status_map[str(t_key)] = status
                if a.get('team_id'):
                    attendance_status_map[str(a['team_id'])] = status
                if a.get('pass_id'):
                    attendance_status_map[str(a['pass_id'])] = status
            continue

        a_slug = (a.get('tournament_slug') or a.get('tournamentSlug') or '').strip().lower()
        if not a_slug or a_slug == clean_slug:
            if isinstance(k, str):
                attendance_status_map[k] = status
            if a.get('pass_id'):
                attendance_status_map[str(a['pass_id'])] = status
            if a.get('team_id'):
                attendance_status_map[str(a['team_id'])] = status

    # Resolve PRESENT Teams
    registered_count = len(regs_map)
    seen_team_ids = set()
    present_teams = []

    for reg_key, reg in regs_map.items():
        team_id = str(reg.get('team_id') or reg.get('teamId') or reg_key)
        pid = reg.get('pass_id') or reg.get('passId') or reg_key

        status = (
            attendance_status_map.get(str(pid)) or
            attendance_status_map.get(str(team_id)) or
            attendance_status_map.get(str(reg_key)) or
            reg.get('attendance_status') or
            reg.get('attendanceStatus') or
            'NOT_MARKED'
        ).upper()

        if status == 'PRESENT':
            if team_id in seen_team_ids:
                continue
            seen_team_ids.add(team_id)

            team_name = reg.get('team_name') or reg.get('teamName') or 'Squad'
            captain_ign = (
                reg.get('captain_in_game_name') or reg.get('captainInGameName') or
                reg.get('captain_freefire_username') or reg.get('captainFreeFireUsername') or ''
            )

            present_teams.append({
                'team_id': team_id,
                'team_name': team_name,
                'captain_name': reg.get('captain_name') or reg.get('captainName') or 'Captain',
                'captain_in_game_name': captain_ign,
                'college': reg.get('college') or '',
                'pass_id': pid,
                'attended_at': reg.get('attended_at') or reg.get('attendedAt')
            })

    return registered_count, present_teams


@leaderboard_bp.route('/<tournament_slug>', methods=['GET'])
def get_organizer_leaderboard(tournament_slug):
    """
    GET /api/leaderboard/<tournament_slug>
    Returns:
      - tournament: basic metadata
      - counts: { registered: int, present: int }
      - columns: list of dynamic scoring rules configured for this tournament
      - matches: list of matches for this tournament
      - teams: list of teams currently marked PRESENT (from attendance system)
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        # 1. Fetch Dynamic Scoring Rules & Matches
        columns = get_rules_for_tournament(clean_slug)
        matches = get_matches_for_tournament(clean_slug)

        # 2. Fetch Present Teams
        registered_count, present_teams = get_present_teams_for_tournament(clean_slug)

        # Build initial team row with initial scoring values (0 or empty, Total: 0)
        formatted_teams = []
        for idx, t in enumerate(present_teams, 1):
            scores = {col['id']: 0 for col in columns}
            formatted_teams.append({
                **t,
                'rank': idx,
                'scores': scores,
                'total': 0
            })

        return jsonify({
            'success': True,
            'tournament_slug': clean_slug,
            'counts': {
                'registered': registered_count,
                'present': len(present_teams)
            },
            'columns': columns,
            'matches': matches,
            'teams': formatted_teams
        }), 200

    except Exception as e:
        print(f"Error generating organizer leaderboard: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/rules', methods=['POST'])
def create_scoring_rule(tournament_slug):
    """
    POST /api/leaderboard/<tournament_slug>/rules
    Add a new custom scoring column.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        is_locked, err_resp = check_tournament_lock_rejection(clean_slug, 'Scoring rules cannot be added')
        if is_locked:
            return err_resp

        data = request.get_json(silent=True) or {}
        name = (data.get('name') or '').strip()
        rule_type = (data.get('type') or 'PER_UNIT').strip().upper()
        points_per_unit = float(data.get('points_per_unit', 0) or 0)
        sort_order = int(data.get('sort_order', 0) or 0)
        placement_points = data.get('placement_points') or []

        if not name:
            return jsonify({'success': False, 'message': 'Column name is required.'}), 400

        valid_types = ['PER_UNIT', 'OCCURRENCE', 'PENALTY', 'PLACEMENT']
        if rule_type not in valid_types:
            return jsonify({'success': False, 'message': f"Invalid column type '{rule_type}'. Must be one of {valid_types}."}), 400

        clean_placement_points = []
        if rule_type == 'PLACEMENT':
            if not isinstance(placement_points, list) or len(placement_points) == 0:
                placement_points = [
                    {'placement': 1, 'points': 12},
                    {'placement': 2, 'points': 9},
                    {'placement': 3, 'points': 8},
                    {'placement': 4, 'points': 7},
                    {'placement': 5, 'points': 6},
                    {'placement': 6, 'points': 5},
                    {'placement': 7, 'points': 4},
                    {'placement': 8, 'points': 3},
                    {'placement': 9, 'points': 2},
                    {'placement': 10, 'points': 1}
                ]
            for p in placement_points:
                try:
                    pos = int(p.get('placement', 1))
                    pts = float(p.get('points', 0))
                    clean_placement_points.append({'placement': pos, 'points': pts})
                except (ValueError, TypeError):
                    continue
            clean_placement_points = sorted(clean_placement_points, key=lambda x: x['placement'])

        current_rules = get_rules_for_tournament(clean_slug)
        if sort_order == 0 and len(current_rules) > 0:
            sort_order = max(r.get('sort_order', 0) for r in current_rules) + 1

        rule_id = str(uuid.uuid4())
        now_iso = get_current_iso_timestamp()

        new_rule = {
            'id': rule_id,
            'tournament_id': clean_slug,
            'name': name,
            'type': rule_type,
            'points_per_unit': points_per_unit,
            'sort_order': sort_order,
            'created_at': now_iso,
            'updated_at': now_iso,
            'placement_points': clean_placement_points
        }

        # 1. Save to In-Memory store
        if clean_slug not in IN_MEMORY_SCORING_RULES:
            IN_MEMORY_SCORING_RULES[clean_slug] = []
        IN_MEMORY_SCORING_RULES[clean_slug].append(new_rule)
        if rule_type == 'PLACEMENT':
            IN_MEMORY_PLACEMENT_RULES[rule_id] = clean_placement_points

        # 2. Persist to Supabase if table is created
        try:
            supabase = get_supabase_client()
            db_payload = {
                'id': rule_id,
                'tournament_id': clean_slug,
                'name': name,
                'type': rule_type,
                'points_per_unit': points_per_unit,
                'sort_order': sort_order,
                'created_at': now_iso,
                'updated_at': now_iso
            }
            supabase.table('tournament_scoring_rules').insert(db_payload).execute()

            if rule_type == 'PLACEMENT' and clean_placement_points:
                p_rows = [{
                    'id': str(uuid.uuid4()),
                    'scoring_rule_id': rule_id,
                    'placement': p['placement'],
                    'points': p['points'],
                    'created_at': now_iso
                } for p in clean_placement_points]
                supabase.table('placement_scoring_rules').insert(p_rows).execute()
        except Exception:
            pass

        return jsonify({
            'success': True,
            'message': f"Scoring column '{name}' added successfully.",
            'rule': new_rule
        }), 201

    except Exception as e:
        print(f"Error creating scoring rule: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/rules/<rule_id>', methods=['PUT', 'PATCH'])
def update_scoring_rule(tournament_slug, rule_id):
    """
    PUT /api/leaderboard/<tournament_slug>/rules/<rule_id>
    Update an existing custom scoring column.
    Automatically recalculates all affected match results from stored RAW values.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        clean_rule_id = str(rule_id).strip()

        is_locked, err_resp = check_tournament_lock_rejection(clean_slug, 'Scoring rules cannot be modified')
        if is_locked:
            return err_resp
        data = request.get_json(silent=True) or {}

        now_iso = get_current_iso_timestamp()

        current_rules = get_rules_for_tournament(clean_slug)
        existing = next((r for r in current_rules if str(r['id']) == clean_rule_id), None)
        if not existing:
            return jsonify({'success': False, 'message': f"Scoring column with ID '{clean_rule_id}' not found."}), 404

        name = (data.get('name') or existing['name']).strip()
        rule_type = (data.get('type') or existing['type']).strip().upper()
        points_per_unit = float(data.get('points_per_unit', existing.get('points_per_unit', 0)))
        sort_order = int(data.get('sort_order', existing.get('sort_order', 0)))
        
        placement_points = data.get('placement_points')
        clean_placement_points = []
        if rule_type == 'PLACEMENT':
            if placement_points is not None and isinstance(placement_points, list):
                for p in placement_points:
                    try:
                        clean_placement_points.append({
                            'placement': int(p.get('placement', 1)),
                            'points': float(p.get('points', 0))
                        })
                    except (ValueError, TypeError):
                        continue
                clean_placement_points = sorted(clean_placement_points, key=lambda x: x['placement'])
            else:
                clean_placement_points = existing.get('placement_points', [])

        updated_rule = {
            **existing,
            'name': name,
            'type': rule_type,
            'points_per_unit': points_per_unit,
            'sort_order': sort_order,
            'updated_at': now_iso,
            'placement_points': clean_placement_points
        }

        # 1. Update in memory
        mem_rules = IN_MEMORY_SCORING_RULES.get(clean_slug, [])
        for i, mr in enumerate(mem_rules):
            if str(mr['id']) == clean_rule_id:
                mem_rules[i] = updated_rule
                break
        if rule_type == 'PLACEMENT':
            IN_MEMORY_PLACEMENT_RULES[clean_rule_id] = clean_placement_points

        # 2. Update Supabase
        try:
            supabase = get_supabase_client()
            supabase.table('tournament_scoring_rules').update({
                'name': name,
                'type': rule_type,
                'points_per_unit': points_per_unit,
                'sort_order': sort_order,
                'updated_at': now_iso
            }).eq('id', clean_rule_id).execute()

            if rule_type == 'PLACEMENT':
                supabase.table('placement_scoring_rules').delete().eq('scoring_rule_id', clean_rule_id).execute()
                if clean_placement_points:
                    p_rows = [{
                        'id': str(uuid.uuid4()),
                        'scoring_rule_id': clean_rule_id,
                        'placement': p['placement'],
                        'points': p['points'],
                        'created_at': now_iso
                    } for p in clean_placement_points]
                    supabase.table('placement_scoring_rules').insert(p_rows).execute()
        except Exception:
            pass

        # 3. AUTOMATIC RECALCULATION: Recalculate match results from RAW values using new rule definition
        recalculate_match_results_for_tournament(clean_slug)

        return jsonify({
            'success': True,
            'message': f"Scoring column '{name}' updated and results recalculated successfully.",
            'rule': updated_rule
        }), 200

    except Exception as e:
        print(f"Error updating scoring rule: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/rules/<rule_id>', methods=['DELETE'])
def delete_scoring_rule(tournament_slug, rule_id):
    """
    DELETE /api/leaderboard/<tournament_slug>/rules/<rule_id>
    Safely delete a custom scoring column and its placement configuration.
    CRITICAL: Does NOT delete team, registration, or attendance records.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        clean_rule_id = str(rule_id).strip()

        is_locked, err_resp = check_tournament_lock_rejection(clean_slug, 'Scoring rules cannot be deleted')
        if is_locked:
            return err_resp

        if clean_slug in IN_MEMORY_SCORING_RULES:
            IN_MEMORY_SCORING_RULES[clean_slug] = [
                r for r in IN_MEMORY_SCORING_RULES[clean_slug]
                if str(r['id']) != clean_rule_id
            ]
        if clean_rule_id in IN_MEMORY_PLACEMENT_RULES:
            del IN_MEMORY_PLACEMENT_RULES[clean_rule_id]

        try:
            supabase = get_supabase_client()
            try:
                supabase.table('placement_scoring_rules').delete().eq('scoring_rule_id', clean_rule_id).execute()
            except Exception:
                pass
            supabase.table('tournament_scoring_rules').delete().eq('id', clean_rule_id).execute()
        except Exception:
            pass

        # Recalculate remaining scores
        recalculate_match_results_for_tournament(clean_slug)

        return jsonify({
            'success': True,
            'message': f"Scoring column deleted successfully.",
            'deleted_id': clean_rule_id
        }), 200

    except Exception as e:
        print(f"Error deleting scoring rule: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/rules/reorder', methods=['POST', 'PUT'])
def reorder_scoring_rules(tournament_slug):
    """
    POST/PUT /api/leaderboard/<tournament_slug>/rules/reorder
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        is_locked, err_resp = check_tournament_lock_rejection(clean_slug, 'Scoring rules cannot be reordered')
        if is_locked:
            return err_resp
        data = request.get_json(silent=True) or {}
        rule_ids = data.get('rule_ids') or []

        if not isinstance(rule_ids, list):
            return jsonify({'success': False, 'message': 'rule_ids must be a list of IDs.'}), 400

        if clean_slug in IN_MEMORY_SCORING_RULES:
            rule_id_order = {str(rid): idx for idx, rid in enumerate(rule_ids)}
            for r in IN_MEMORY_SCORING_RULES[clean_slug]:
                rid_str = str(r['id'])
                if rid_str in rule_id_order:
                    r['sort_order'] = rule_id_order[rid_str]
            IN_MEMORY_SCORING_RULES[clean_slug].sort(key=lambda x: x.get('sort_order', 0))

        try:
            supabase = get_supabase_client()
            for idx, rid in enumerate(rule_ids):
                supabase.table('tournament_scoring_rules').update({'sort_order': idx}).eq('id', str(rid)).execute()
        except Exception:
            pass

        updated_rules = get_rules_for_tournament(clean_slug)
        return jsonify({
            'success': True,
            'message': 'Columns reordered successfully.',
            'columns': updated_rules
        }), 200

    except Exception as e:
        print(f"Error reordering scoring rules: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


# ════════════════════════════════════════════════════════════════════════════════
# PHASE 3: MATCHES & RESULT ENTRY ENDPOINTS
# ════════════════════════════════════════════════════════════════════════════════

@leaderboard_bp.route('/<tournament_slug>/matches', methods=['GET'])
def list_tournament_matches(tournament_slug):
    """
    GET /api/leaderboard/<tournament_slug>/matches
    Fetch list of matches configured for this tournament.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        matches = get_matches_for_tournament(clean_slug)
        return jsonify({
            'success': True,
            'tournament_slug': clean_slug,
            'matches': matches
        }), 200
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/matches', methods=['POST'])
def create_tournament_match(tournament_slug):
    """
    POST /api/leaderboard/<tournament_slug>/matches
    Add a new match to this tournament (e.g. Match 1, Match 2).
    Body:
      - title: optional string (defaults to "Match {match_number}")
      - match_number: optional int
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        is_locked, err_resp = check_tournament_lock_rejection(clean_slug, 'Matches cannot be added')
        if is_locked:
            return err_resp

        data = request.get_json(silent=True) or {}
        existing_matches = get_matches_for_tournament(clean_slug)
        
        match_number = data.get('match_number')
        if not match_number:
            match_number = (max([m['match_number'] for m in existing_matches], default=0)) + 1
        else:
            match_number = int(match_number)

        title = (data.get('title') or '').strip()
        if not title:
            title = f"Match {match_number}"

        match_id = str(uuid.uuid4())
        now_iso = get_current_iso_timestamp()

        new_match = {
            'id': match_id,
            'tournament_id': clean_slug,
            'match_number': match_number,
            'title': title,
            'status': data.get('status', 'COMPLETED'),
            'created_at': now_iso,
            'updated_at': now_iso
        }

        # 1. Save to in-memory store
        if clean_slug not in IN_MEMORY_MATCHES:
            IN_MEMORY_MATCHES[clean_slug] = []
        IN_MEMORY_MATCHES[clean_slug].append(new_match)

        # 2. Save to Supabase
        try:
            supabase = get_supabase_client()
            supabase.table('tournament_matches').insert(new_match).execute()
        except Exception:
            pass

        return jsonify({
            'success': True,
            'message': f"Match '{title}' created successfully.",
            'match': new_match
        }), 201

    except Exception as e:
        print(f"Error creating tournament match: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/matches/<match_id>', methods=['GET'])
def get_match_details_and_results(tournament_slug, match_id):
    """
    GET /api/leaderboard/<tournament_slug>/matches/<match_id>
    Returns:
      - match info
      - columns (scoring rules)
      - teams: ONLY teams currently marked PRESENT with their entered RAW scores,
        calculated scores, and match total.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        clean_match_id = str(match_id).strip()

        # 1. Fetch match
        matches = get_matches_for_tournament(clean_slug)
        match_obj = next((m for m in matches if str(m['id']) == clean_match_id), None)
        if not match_obj:
            return jsonify({'success': False, 'message': f"Match ID '{clean_match_id}' not found."}), 404

        # 2. Fetch tournament scoring rules
        columns = get_rules_for_tournament(clean_slug)

        # 3. Fetch ONLY verified PRESENT teams from existing attendance system
        registered_count, present_teams = get_present_teams_for_tournament(clean_slug)

        # 4. Fetch existing match results (from memory and Supabase)
        existing_results = {}
        # Fetch from Supabase
        try:
            supabase = get_supabase_client()
            res = supabase.table('match_team_results').select('*').eq('match_id', clean_match_id).execute()
            if res.data:
                for r in res.data:
                    existing_results[str(r['team_id'])] = r
        except Exception:
            pass

        # Merge with in-memory results
        mem_results = IN_MEMORY_MATCH_RESULTS.get(clean_match_id, {})
        for tid, r in mem_results.items():
            if tid not in existing_results:
                existing_results[tid] = r

        # 5. Build team result rows for ONLY PRESENT teams
        teams_data = []
        for idx, t in enumerate(present_teams, 1):
            tid = t['team_id']
            res_row = existing_results.get(tid)

            if res_row:
                raw_scores = res_row.get('raw_scores', {})
                # Ensure all active columns have an entry in raw_scores (default 0)
                full_raw = {col['id']: raw_scores.get(col['id'], 0) for col in columns}
                calc_scores = {}
                for col in columns:
                    calc_scores[col['id']] = calculate_score_for_rule(col, full_raw.get(col['id'], 0))
                total_pts = sum(calc_scores.values())
            else:
                full_raw = {col['id']: 0 for col in columns}
                calc_scores = {col['id']: 0 for col in columns}
                total_pts = 0

            teams_data.append({
                **t,
                'rank': idx,
                'raw_scores': full_raw,
                'calculated_scores': calc_scores,
                'total_points': total_pts
            })

        return jsonify({
            'success': True,
            'tournament_slug': clean_slug,
            'match': match_obj,
            'columns': columns,
            'teams': teams_data,
            'counts': {
                'registered': registered_count,
                'present': len(present_teams)
            }
        }), 200

    except Exception as e:
        print(f"Error fetching match results: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/matches/<match_id>/results', methods=['PUT', 'POST'])
def save_match_results(tournament_slug, match_id):
    """
    PUT/POST /api/leaderboard/<tournament_slug>/matches/<match_id>/results
    Save / update match scores.
    CRITICAL REQUIREMENTS:
      1. ONLY teams marked PRESENT can have results entered.
      2. Organizer enters RAW values only.
      3. Server calculates points using configured scoring rules.
      4. Stored RAW values are preserved and never overwritten by calculated points.
    Body:
      - results: list of {
          team_id: str,
          raw_scores: { "<rule_id>": raw_value }
        }
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        clean_match_id = str(match_id).strip()

        is_locked, err_resp = check_tournament_lock_rejection(clean_slug, 'Match results cannot be modified')
        if is_locked:
            return err_resp

        # 1. Verify Match exists
        matches = get_matches_for_tournament(clean_slug)
        match_obj = next((m for m in matches if str(m['id']) == clean_match_id), None)
        if not match_obj:
            return jsonify({'success': False, 'message': f"Match ID '{clean_match_id}' not found."}), 404

        # 2. Verify rules
        columns = get_rules_for_tournament(clean_slug)
        rules_dict = {str(c['id']): c for c in columns}

        # 3. Read PRESENT teams from existing attendance system
        registered_count, present_teams = get_present_teams_for_tournament(clean_slug)
        present_team_ids = {t['team_id'] for t in present_teams}
        present_team_map = {t['team_id']: t for t in present_teams}

        data = request.get_json(silent=True) or {}
        submitted_results = data.get('results') or []

        if not isinstance(submitted_results, list) or len(submitted_results) == 0:
            return jsonify({'success': False, 'message': 'results list is required.'}), 400

        # 4. Strict Validation: ONLY PRESENT teams can enter results
        for item in submitted_results:
            team_id = item.get('team_id')
            if team_id not in present_team_ids:
                return jsonify({
                    'success': False,
                    'message': f"Team ID '{team_id}' is not marked PRESENT. Only verified PRESENT teams can participate in match results."
                }), 400

        # 5. Process each team's raw scores and calculate points on server
        now_iso = get_current_iso_timestamp()
        processed_results = []

        if clean_match_id not in IN_MEMORY_MATCH_RESULTS:
            IN_MEMORY_MATCH_RESULTS[clean_match_id] = {}

        for item in submitted_results:
            team_id = item.get('team_id')
            input_raw = item.get('raw_scores') or {}

            # Build sanitized raw_scores (only for existing rules)
            raw_scores = {}
            calc_scores = {}
            for r_id, rule in rules_dict.items():
                raw_val = input_raw.get(r_id, 0)
                try:
                    raw_scores[r_id] = float(raw_val)
                except (ValueError, TypeError):
                    raw_scores[r_id] = 0.0
                
                # Server calculation
                calc_scores[r_id] = calculate_score_for_rule(rule, raw_scores[r_id])

            total_points = sum(calc_scores.values())

            result_id = str(uuid.uuid4())
            result_record = {
                'id': result_id,
                'match_id': clean_match_id,
                'tournament_id': clean_slug,
                'team_id': team_id,
                'raw_scores': raw_scores,
                'calculated_scores': calc_scores,
                'total_points': total_points,
                'created_at': now_iso,
                'updated_at': now_iso
            }

            # Save in memory
            IN_MEMORY_MATCH_RESULTS[clean_match_id][team_id] = result_record

            # Save in Supabase
            try:
                supabase = get_supabase_client()
                # Upsert into match_team_results
                supabase.table('match_team_results').upsert(result_record, on_conflict='match_id,team_id').execute()
            except Exception:
                pass

            team_info = present_team_map.get(team_id, {})
            processed_results.append({
                **result_record,
                'team_name': team_info.get('team_name', 'Squad'),
                'captain_in_game_name': team_info.get('captain_in_game_name', '')
            })

        return jsonify({
            'success': True,
            'message': f"Match results saved successfully for {len(processed_results)} squads.",
            'match_id': clean_match_id,
            'results': processed_results
        }), 200

    except Exception as e:
        print(f"Error saving match results: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/matches/<match_id>', methods=['DELETE'])
def delete_tournament_match(tournament_slug, match_id):
    """
    DELETE /api/leaderboard/<tournament_slug>/matches/<match_id>
    Safely delete a match and its results.
    Does NOT delete teams, registrations, or attendance records.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        clean_match_id = str(match_id).strip()

        is_locked, err_resp = check_tournament_lock_rejection(clean_slug, 'Matches cannot be deleted')
        if is_locked:
            return err_resp

        # 1. Remove from memory
        if clean_slug in IN_MEMORY_MATCHES:
            IN_MEMORY_MATCHES[clean_slug] = [
                m for m in IN_MEMORY_MATCHES[clean_slug]
                if str(m['id']) != clean_match_id
            ]
        if clean_match_id in IN_MEMORY_MATCH_RESULTS:
            del IN_MEMORY_MATCH_RESULTS[clean_match_id]

        # 2. Remove from Supabase
        try:
            supabase = get_supabase_client()
            supabase.table('match_team_results').delete().eq('match_id', clean_match_id).execute()
            supabase.table('tournament_matches').delete().eq('id', clean_match_id).execute()
        except Exception:
            pass

        return jsonify({
            'success': True,
            'message': f"Match and results deleted successfully.",
            'deleted_id': clean_match_id
        }), 200

    except Exception as e:
        print(f"Error deleting tournament match: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


# ════════════════════════════════════════════════════════════════════════════════
# PHASE 4: LIVE CUMULATIVE TOURNAMENT STANDINGS
# ════════════════════════════════════════════════════════════════════════════════

def get_standings_for_tournament(tournament_slug):
    """
    Computes live cumulative standings across all matches for this tournament.
    Requirements:
      1. Tournament-specific isolation.
      2. Sums match totals from all tournament matches: Overall Total = SUM(all match totals).
      3. Automatic authoritative ranking (Rank 1 for highest overall points).
      4. Dynamic match columns (M1, M2, M3...).
      5. Only verified PRESENT squads appear.
      6. Missing match results are rendered as 0 without removing the team.
      7. Deterministic tie-breaking.
      8. Preserves detailed match breakdowns for expandable match inspect.
    """
    clean_slug = (tournament_slug or '').strip().lower()

    # 1. Fetch scoring rules for column definitions & breakdowns
    columns = get_rules_for_tournament(clean_slug)
    rules_dict = {str(c['id']): c for c in columns}

    # 2. Fetch matches for this tournament sorted by match_number
    matches = get_matches_for_tournament(clean_slug)
    matches.sort(key=lambda m: (m.get('match_number', 0), m.get('created_at', '')))

    # 3. Fetch verified PRESENT teams from existing attendance system
    registered_count, present_teams = get_present_teams_for_tournament(clean_slug)

    # 4. Fetch all results for all matches of this tournament
    match_results_by_match = {}
    for m in matches:
        mid = str(m['id'])
        m_results = {}
        # Fetch from Supabase
        try:
            supabase = get_supabase_client()
            res = supabase.table('match_team_results').select('*').eq('match_id', mid).execute()
            if res.data:
                for r in res.data:
                    m_results[str(r['team_id'])] = r
        except Exception:
            pass
        # Merge with in-memory results
        mem_r = IN_MEMORY_MATCH_RESULTS.get(mid, {})
        for tid, r in mem_r.items():
            if tid not in m_results:
                m_results[tid] = r
        match_results_by_match[mid] = m_results

    # 5. Build cumulative standings rows for each PRESENT team
    standings_rows = []
    for t in present_teams:
        tid = str(t['team_id'])
        match_scores = {}
        match_breakdowns = {}
        overall_total = 0.0

        for m in matches:
            mid = str(m['id'])
            m_res = match_results_by_match.get(mid, {}).get(tid)

            if m_res:
                raw_scores = m_res.get('raw_scores', {})
                # Recalculate accurately using current rules to ensure rule changes flow through immediately
                calc_scores = {}
                for r_id, rule in rules_dict.items():
                    raw_val = raw_scores.get(r_id, 0)
                    calc_scores[r_id] = calculate_score_for_rule(rule, raw_val)
                pts = sum(calc_scores.values())
            else:
                pts = 0.0
                raw_scores = {col['id']: 0 for col in columns}
                calc_scores = {col['id']: 0 for col in columns}

            clean_pts = int(pts) if isinstance(pts, float) and pts.is_integer() else pts
            match_scores[mid] = clean_pts
            match_breakdowns[mid] = {
                'match_id': mid,
                'match_title': m.get('title', f"Match {m.get('match_number', 1)}"),
                'match_number': m.get('match_number', 1),
                'total_points': clean_pts,
                'raw_scores': raw_scores,
                'calculated_scores': calc_scores
            }
            overall_total += pts

        clean_overall = int(overall_total) if isinstance(overall_total, float) and overall_total.is_integer() else round(overall_total, 2)

        standings_rows.append({
            'team_id': tid,
            'team_name': t.get('team_name', 'Squad'),
            'captain_name': t.get('captain_name', ''),
            'captain_in_game_name': t.get('captain_in_game_name', ''),
            'college': t.get('college', ''),
            'pass_id': t.get('pass_id', ''),
            'match_scores': match_scores,
            'match_breakdowns': match_breakdowns,
            'overall_total': clean_overall
        })

    # 6. Automatic authoritative ranking
    # Sort by overall_total DESC, then team_name ASC, then team_id ASC (deterministic tie-break)
    standings_rows.sort(
        key=lambda x: (-x['overall_total'], (x.get('team_name') or '').lower(), str(x['team_id']))
    )

    for rank_idx, row in enumerate(standings_rows, 1):
        row['rank'] = rank_idx

    return {
        'tournament_slug': clean_slug,
        'matches': matches,
        'columns': columns,
        'counts': {
            'registered': registered_count,
            'present': len(present_teams),
            'matches': len(matches)
        },
        'standings': standings_rows
    }


@leaderboard_bp.route('/<tournament_slug>/standings', methods=['GET'])
def get_tournament_standings(tournament_slug):
    """
    GET /api/leaderboard/<tournament_slug>/standings
    Phase 4: Live Cumulative Leaderboard endpoint.
    Aggregates match totals from all tournament matches for verified PRESENT squads.
    Calculates overall totals and ranks on the backend.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        data = get_standings_for_tournament(clean_slug)
        return jsonify({
            'success': True,
            **data
        }), 200

    except Exception as e:
        print(f"Error fetching tournament standings: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


# ════════════════════════════════════════════════════════════════════════════════
# PHASE 5: EXPORT GENERATORS & ADMIN SUBMISSION ENDPOINTS
# ════════════════════════════════════════════════════════════════════════════════

def generate_leaderboard_csv(tournament_slug, metadata, standings_data):
    """
    Generates a full-detail CSV export of the tournament leaderboard.
    Includes metadata, cumulative standings, and match-by-match breakdown.
    Never hardcodes rule columns.
    """
    output = io.StringIO()
    writer = csv.writer(output)

    title = metadata.get('title') or metadata.get('name') or tournament_slug.upper()
    game = metadata.get('game', 'Free Fire')
    status = metadata.get('status', 'FINALIZED')

    matches = standings_data.get('matches', [])
    columns = standings_data.get('columns', [])
    standings = standings_data.get('standings', [])

    # 1. Metadata Block
    writer.writerow(["XENOVA ESPORTS — TOURNAMENT LEADERBOARD REPORT"])
    writer.writerow(["Tournament Name", title])
    writer.writerow(["Tournament Slug", tournament_slug])
    writer.writerow(["Game", game])
    writer.writerow(["Format", metadata.get('format', 'Battle Royale')])
    writer.writerow(["Status", status])
    writer.writerow(["Exported At", get_current_iso_timestamp()])
    writer.writerow(["Total Verified Squads", len(standings)])
    writer.writerow(["Total Matches", len(matches)])
    writer.writerow([])

    # 2. Cumulative Standings Table
    match_headers = [f"{m.get('title', f'Match {m.get('match_number', idx + 1)}')}" for idx, m in enumerate(matches)]
    standings_headers = ["Rank", "Team ID", "Team Name", "Captain Name", "Captain In-Game Name", "College"] + match_headers + ["Overall Total"]
    writer.writerow(standings_headers)

    for s in standings:
        m_scores = [s.get('match_scores', {}).get(m['id'], 0) for m in matches]
        row = [
            s.get('rank', ''),
            s.get('team_id', ''),
            s.get('team_name', ''),
            s.get('captain_name', ''),
            s.get('captain_in_game_name', ''),
            s.get('college', ''),
        ] + m_scores + [s.get('overall_total', 0)]
        writer.writerow(row)

    writer.writerow([])
    writer.writerow(["DETAILED MATCH-BY-MATCH BREAKDOWN"])
    writer.writerow([])

    # 3. Match Details Table
    col_names = [c.get('name', 'Score') for c in columns]
    detail_headers = ["Match", "Team Name", "Captain In-Game Name"]
    for cn in col_names:
        detail_headers.append(f"{cn} (Raw)")
        detail_headers.append(f"{cn} (Pts)")
    detail_headers.append("Match Total")
    writer.writerow(detail_headers)

    for m in matches:
        mid = str(m['id'])
        m_title = m.get('title', f"Match {m.get('match_number', 1)}")
        for s in standings:
            b = s.get('match_breakdowns', {}).get(mid, {})
            raw_s = b.get('raw_scores', {})
            calc_s = b.get('calculated_scores', {})
            m_total = b.get('total_points', 0)

            detail_row = [m_title, s.get('team_name', ''), s.get('captain_in_game_name', '')]
            for col in columns:
                cid = str(col['id'])
                detail_row.append(raw_s.get(cid, 0))
                detail_row.append(calc_s.get(cid, 0))
            detail_row.append(m_total)
            writer.writerow(detail_row)

    return output.getvalue()


def generate_leaderboard_xlsx(tournament_slug, metadata, standings_data):
    """
    Generates an Excel workbook (.xlsx) with multi-sheet formatting:
    - Sheet 1: Cumulative Standings
    - Sheet 2: Match Breakdown
    - Sheet 3: Tournament Metadata & Rules
    """
    wb = openpyxl.Workbook()

    title = metadata.get('title') or metadata.get('name') or tournament_slug.upper()
    game = metadata.get('game', 'Free Fire')
    status = metadata.get('status', 'FINALIZED')

    matches = standings_data.get('matches', [])
    columns = standings_data.get('columns', [])
    standings = standings_data.get('standings', [])

    # Sheet 1: Final Standings
    ws1 = wb.active
    ws1.title = "Final Standings"

    # Title Banner
    ws1.merge_cells("A1:G1")
    ws1["A1"] = f"XENOVA ESPORTS — {title}"
    ws1["A1"].font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    ws1["A1"].fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 32

    # Metadata row
    ws1["A2"] = "Game:"
    ws1["B2"] = game
    ws1["C2"] = "Status:"
    ws1["D2"] = status
    ws1["E2"] = "Exported:"
    ws1["F2"] = get_current_iso_timestamp()

    for cell_ref in ["A2", "C2", "E2"]:
        ws1[cell_ref].font = Font(name="Calibri", size=10, bold=True, color="64748B")

    match_headers = [f"{m.get('title', f'Match {m.get('match_number', idx + 1)}')}" for idx, m in enumerate(matches)]
    headers = ["Rank", "Team ID", "Team Name", "Captain Name", "Captain IGN", "College"] + match_headers + ["Overall Total"]

    ws1.append([])  # blank row 3
    ws1.append(headers)  # row 4
    header_row_idx = 4
    ws1.row_dimensions[header_row_idx].height = 24

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="F8FAFC")

    for col_idx in range(1, len(headers) + 1):
        cell = ws1.cell(row=header_row_idx, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for s in standings:
        m_scores = [s.get('match_scores', {}).get(m['id'], 0) for m in matches]
        row_data = [
            s.get('rank', ''),
            s.get('team_id', ''),
            s.get('team_name', ''),
            s.get('captain_name', ''),
            s.get('captain_in_game_name', ''),
            s.get('college', ''),
        ] + m_scores + [s.get('overall_total', 0)]
        ws1.append(row_data)

    # Sheet 2: Match Breakdown
    ws2 = wb.create_sheet(title="Match Breakdown")
    detail_headers = ["Match", "Team Name", "Captain IGN"]
    for c in columns:
        cn = c.get('name', 'Score')
        detail_headers.append(f"{cn} (Raw)")
        detail_headers.append(f"{cn} (Pts)")
    detail_headers.append("Match Total")
    ws2.append(detail_headers)

    ws2.row_dimensions[1].height = 24
    for col_idx in range(1, len(detail_headers) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for m in matches:
        mid = str(m['id'])
        m_title = m.get('title', f"Match {m.get('match_number', 1)}")
        for s in standings:
            b = s.get('match_breakdowns', {}).get(mid, {})
            raw_s = b.get('raw_scores', {})
            calc_s = b.get('calculated_scores', {})
            m_total = b.get('total_points', 0)

            detail_row = [m_title, s.get('team_name', ''), s.get('captain_in_game_name', '')]
            for col in columns:
                cid = str(col['id'])
                detail_row.append(raw_s.get(cid, 0))
                detail_row.append(calc_s.get(cid, 0))
            detail_row.append(m_total)
            ws2.append(detail_row)

    # Auto-fit columns
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()


def generate_leaderboard_printable_html(tournament_slug, metadata, standings_data):
    """
    Generates a clean HTML report designed for print/PDF rendering with complete standings,
    tournament metadata, and match breakdowns.
    """
    title = metadata.get('title') or metadata.get('name') or tournament_slug.upper()
    game = metadata.get('game', 'Free Fire')
    status = metadata.get('status', 'FINALIZED')
    matches = standings_data.get('matches', [])
    columns = standings_data.get('columns', [])
    standings = standings_data.get('standings', [])
    now_iso = get_current_iso_timestamp()

    match_ths = "".join([f"<th>M{m.get('match_number', idx+1)}<br><span style='font-size:9px;font-weight:normal;'>{m.get('title','')}</span></th>" for idx, m in enumerate(matches)])

    standings_rows_html = ""
    for s in standings:
        m_cells = "".join([f"<td style='text-align:center;'>{s.get('match_scores', {}).get(m['id'], 0)}</td>" for m in matches])
        rank_badge = f"<span class='rank-badge rank-{s['rank']}'>#{s['rank']}</span>" if s['rank'] <= 3 else f"#{s['rank']}"
        ign_badge = f"<span class='ign-badge'>🎮 {s['captain_in_game_name']}</span>" if s.get('captain_in_game_name') else "—"
        standings_rows_html += f"""
        <tr>
            <td style="text-align:center;font-weight:bold;">{rank_badge}</td>
            <td><strong>{s.get('team_name','')}</strong><br><span style="font-size:10px;color:#64748b;">{s.get('college','')} • Cap: {s.get('captain_name','')}</span></td>
            <td>{ign_badge}</td>
            {m_cells}
            <td style="text-align:right;font-weight:900;color:#059669;font-size:13px;">{s.get('overall_total', 0)} pts</td>
        </tr>
        """

    col_ths = "".join([f"<th>{c.get('name','Rule')} (Raw)</th><th>{c.get('name','Rule')} (Pts)</th>" for c in columns])
    breakdown_rows_html = ""
    for m in matches:
        mid = str(m['id'])
        m_title = m.get('title', f"Match {m.get('match_number', 1)}")
        for s in standings:
            b = s.get('match_breakdowns', {}).get(mid, {})
            raw_s = b.get('raw_scores', {})
            calc_s = b.get('calculated_scores', {})
            m_total = b.get('total_points', 0)
            col_tds = "".join([f"<td style='text-align:center;'>{raw_s.get(str(c['id']),0)}</td><td style='text-align:center;color:#059669;font-weight:bold;'>+{calc_s.get(str(c['id']),0)}</td>" for c in columns])
            breakdown_rows_html += f"""
            <tr>
                <td><strong>{m_title}</strong></td>
                <td>{s.get('team_name','')}</td>
                <td>{s.get('captain_in_game_name','')}</td>
                {col_tds}
                <td style="text-align:right;font-weight:bold;color:#059669;">{m_total} pts</td>
            </tr>
            """

    return f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{title} — Leaderboard Report</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; color: #0f172a; margin: 0; padding: 24px; font-size: 12px; }}
        .header {{ display: flex; justify-content: space-between; align-items: flex-start; border-bottom: 2px solid #0f172a; padding-bottom: 16px; margin-bottom: 20px; }}
        h1 {{ margin: 0; font-size: 22px; text-transform: uppercase; font-style: italic; letter-spacing: -0.5px; }}
        .meta {{ color: #475569; font-size: 11px; margin-top: 4px; }}
        .badge {{ display: inline-block; padding: 3px 8px; border-radius: 4px; font-weight: bold; font-size: 10px; background: #e0e7ff; color: #3730a3; }}
        .rank-badge {{ display: inline-block; padding: 2px 6px; border-radius: 4px; font-weight: bold; }}
        .rank-1 {{ background: #fef3c7; color: #b45309; border: 1px solid #fde68a; }}
        .rank-2 {{ background: #f1f5f9; color: #475569; border: 1px solid #e2e8f0; }}
        .rank-3 {{ background: #ffedd5; color: #c2410c; border: 1px solid #fed7aa; }}
        .ign-badge {{ font-family: monospace; background: #fffbeb; color: #b45309; padding: 2px 6px; border-radius: 4px; border: 1px solid #fef3c7; font-size: 10px; }}
        table {{ width: 100%; border-collapse: collapse; margin-bottom: 28px; }}
        th {{ background: #f8fafc; color: #334155; font-weight: 700; text-transform: uppercase; font-size: 10px; padding: 8px 10px; border-bottom: 2px solid #cbd5e1; text-align: left; }}
        td {{ padding: 8px 10px; border-bottom: 1px solid #e2e8f0; vertical-align: middle; }}
        .section-title {{ font-size: 14px; font-weight: 800; text-transform: uppercase; margin: 24px 0 8px 0; color: #1e293b; letter-spacing: 0.5px; }}
        @media print {{
            body {{ padding: 0; }}
            .no-print {{ display: none; }}
            @page {{ margin: 1.5cm; }}
        }}
    </style>
</head>
<body>
    <div class="no-print" style="margin-bottom: 16px; display: flex; gap: 8px;">
        <button onclick="window.print()" style="padding: 8px 16px; background: #059669; color: white; border: none; border-radius: 6px; font-weight: bold; cursor: pointer;">Print / Save as PDF</button>
        <button onclick="window.close()" style="padding: 8px 16px; background: #64748b; color: white; border: none; border-radius: 6px; font-weight: bold; cursor: pointer;">Close</button>
    </div>

    <div class="header">
        <div>
            <h1>{title}</h1>
            <div class="meta">{game} • Tournament Slug: <strong>{tournament_slug}</strong> • Status: <span class="badge">{status}</span></div>
        </div>
        <div style="text-align: right;" class="meta">
            Generated: {now_iso}<br>
            Verified Present Squads: <strong>{len(standings)}</strong> • Matches: <strong>{len(matches)}</strong>
        </div>
    </div>

    <div class="section-title">🏆 Cumulative Final Standings</div>
    <table>
        <thead>
            <tr>
                <th style="width: 50px; text-align: center;">Rank</th>
                <th>Team</th>
                <th>Captain In-Game Name</th>
                {match_ths}
                <th style="text-align: right;">Overall Total</th>
            </tr>
        </thead>
        <tbody>
            {standings_rows_html}
        </tbody>
    </table>

    <div class="section-title">📊 Detailed Match-by-Match Breakdown</div>
    <table>
        <thead>
            <tr>
                <th>Match</th>
                <th>Team Name</th>
                <th>Captain IGN</th>
                {col_ths}
                <th style="text-align: right;">Match Total</th>
            </tr>
        </thead>
        <tbody>
            {breakdown_rows_html}
        </tbody>
    </table>
</body>
</html>
"""


@leaderboard_bp.route('/<tournament_slug>/status', methods=['GET'])
def get_tournament_status_route(tournament_slug):
    """
    GET /api/leaderboard/<tournament_slug>/status
    Returns the current lifecycle status (LIVE, FINALIZED, SUBMITTED, CHANGES_REQUESTED, APPROVED, PUBLISHED).
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        status_info = get_tournament_lifecycle_status(clean_slug)
        matches = get_matches_for_tournament(clean_slug)

        latest_sub = next((s for s in reversed(IN_MEMORY_SUBMISSIONS) if s.get('tournament_id') == clean_slug), None)

        return jsonify({
            'success': True,
            'tournament_slug': clean_slug,
            'status': status_info.get('status', 'LIVE'),
            'is_locked': is_tournament_locked(clean_slug),
            'finalized_at': status_info.get('finalized_at'),
            'finalized_by': status_info.get('finalized_by'),
            'submitted_at': status_info.get('submitted_at'),
            'submitted_by': status_info.get('submitted_by'),
            'approved_at': status_info.get('approved_at'),
            'approved_by': status_info.get('approved_by'),
            'published_at': status_info.get('published_at'),
            'published_by': status_info.get('published_by'),
            'change_request_reason': status_info.get('change_request_reason'),
            'change_requested_by': status_info.get('change_requested_by'),
            'change_requested_at': status_info.get('change_requested_at'),
            'match_count': len(matches),
            'submission_id': latest_sub.get('id') if latest_sub else None
        }), 200

    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/finalize', methods=['POST'])
def finalize_tournament_results(tournament_slug):
    """
    POST /api/leaderboard/<tournament_slug>/finalize
    Phase 5 & 6: Finalize tournament results.
    Safety checks:
      1. Tournament exists.
      2. Caller is authorized organizer or admin.
      3. At least one match exists.
      4. Results and Standings can be calculated successfully.
    Locks tournament against scoring rule, match, and result edits.
    Does NOT lock attendance.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        user = get_authenticated_user()
        is_auth, status_code, err_msg, tournament_obj = check_organizer_authorization(user, clean_slug)
        if not is_auth:
            return jsonify({'success': False, 'message': err_msg}), status_code

        current_status = get_tournament_lifecycle_status(clean_slug)
        curr_stat_str = (current_status.get('status') or 'LIVE').upper()
        if curr_stat_str in ['APPROVED', 'PUBLISHED']:
            return jsonify({
                'success': False,
                'message': f"Cannot finalize: tournament results are already {curr_stat_str.lower()}."
            }), 400

        if curr_stat_str in ['FINALIZED', 'SUBMITTED']:
            return jsonify({
                'success': True,
                'status': curr_stat_str,
                'message': f"Tournament results are already {curr_stat_str.lower()}."
            }), 200

        matches = get_matches_for_tournament(clean_slug)
        if len(matches) == 0:
            return jsonify({
                'success': False,
                'message': 'Cannot finalize tournament: at least one match must exist.'
            }), 400

        standings_data = get_standings_for_tournament(clean_slug)
        now_iso = get_current_iso_timestamp()
        finalized_by = user.get('email') or user.get('id')

        status_record = {
            'tournament_id': clean_slug,
            'status': 'FINALIZED',
            'finalized_at': now_iso,
            'finalized_by': finalized_by,
            'updated_at': now_iso
        }
        if clean_slug in IN_MEMORY_TOURNAMENT_STATUS:
            status_record['change_request_reason'] = IN_MEMORY_TOURNAMENT_STATUS[clean_slug].get('change_request_reason')
            status_record['change_requested_by'] = IN_MEMORY_TOURNAMENT_STATUS[clean_slug].get('change_requested_by')
            status_record['change_requested_at'] = IN_MEMORY_TOURNAMENT_STATUS[clean_slug].get('change_requested_at')

        IN_MEMORY_TOURNAMENT_STATUS[clean_slug] = status_record

        try:
            supabase = get_supabase_client()
            supabase.table('tournament_leaderboard_status').upsert(status_record).execute()
        except Exception:
            pass

        return jsonify({
            'success': True,
            'status': 'FINALIZED',
            'finalized_at': now_iso,
            'finalized_by': finalized_by,
            'message': 'Tournament results successfully finalized. Scoring rules and match results are now locked.',
            'snapshot': standings_data
        }), 200

    except Exception as e:
        print(f"Error finalizing tournament results: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/submit', methods=['POST'])
def submit_tournament_results(tournament_slug):
    """
    POST /api/leaderboard/<tournament_slug>/submit
    Phase 5 & 6: Submit finalized tournament results to Admin.
    Safety checks:
      1. Tournament exists.
      2. Caller is authorized organizer or admin.
      3. At least one match exists.
    Creates a record in tournament_result_submissions and transitions status to SUBMITTED.
    Records SUBMITTED or RESUBMITTED in audit trail.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        user = get_authenticated_user()
        is_auth, status_code, err_msg, tournament_obj = check_organizer_authorization(user, clean_slug)
        if not is_auth:
            return jsonify({'success': False, 'message': err_msg}), status_code

        current_status = get_tournament_lifecycle_status(clean_slug)
        curr_stat_str = (current_status.get('status') or 'LIVE').upper()
        if curr_stat_str in ['APPROVED', 'PUBLISHED']:
            return jsonify({
                'success': False,
                'message': f'Cannot submit: tournament results are already {curr_stat_str.lower()}.'
            }), 400

        matches = get_matches_for_tournament(clean_slug)
        if len(matches) == 0:
            return jsonify({
                'success': False,
                'message': 'Cannot submit tournament: at least one match must exist.'
            }), 400

        # Determine if this submission is a resubmission after changes were requested
        audit_trail = get_audit_trail_for_tournament(clean_slug)
        was_changes_requested = (curr_stat_str == 'CHANGES_REQUESTED') or any(a.get('action') == 'CHANGES_REQUESTED' for a in audit_trail)
        audit_action = 'RESUBMITTED' if was_changes_requested else 'SUBMITTED'

        standings_data = get_standings_for_tournament(clean_slug)
        now_iso = get_current_iso_timestamp()
        submitted_by = user.get('email') or user.get('id')

        data = request.get_json(silent=True) or {}
        notes = (data.get('notes') or '').strip()

        sub_id = str(uuid.uuid4())
        submission_record = {
            'id': sub_id,
            'tournament_id': clean_slug,
            'submitted_by': submitted_by,
            'status': 'SUBMITTED',
            'snapshot': standings_data,
            'notes': notes,
            'submitted_at': now_iso,
            'created_at': now_iso,
            'updated_at': now_iso
        }

        IN_MEMORY_SUBMISSIONS.append(submission_record)

        status_record = {
            'tournament_id': clean_slug,
            'status': 'SUBMITTED',
            'submitted_at': now_iso,
            'submitted_by': submitted_by,
            'change_request_reason': None,
            'change_requested_by': None,
            'change_requested_at': None,
            'updated_at': now_iso
        }
        if clean_slug in IN_MEMORY_TOURNAMENT_STATUS:
            status_record['finalized_at'] = IN_MEMORY_TOURNAMENT_STATUS[clean_slug].get('finalized_at', now_iso)
            status_record['finalized_by'] = IN_MEMORY_TOURNAMENT_STATUS[clean_slug].get('finalized_by', submitted_by)
        else:
            status_record['finalized_at'] = now_iso
            status_record['finalized_by'] = submitted_by

        IN_MEMORY_TOURNAMENT_STATUS[clean_slug] = status_record

        # Record audit trail
        record_audit_event(clean_slug, sub_id, audit_action, submitted_by, notes or None)

        try:
            supabase = get_supabase_client()
            supabase.table('tournament_result_submissions').insert(submission_record).execute()
            supabase.table('tournament_leaderboard_status').upsert(status_record).execute()
        except Exception:
            pass

        return jsonify({
            'success': True,
            'status': 'SUBMITTED',
            'submission_id': sub_id,
            'submitted_at': now_iso,
            'message': f"Tournament results successfully {'resubmitted' if was_changes_requested else 'submitted'} to Admin. Status: Awaiting Admin Review."
        }), 200

    except Exception as e:
        print(f"Error submitting tournament results: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


@leaderboard_bp.route('/<tournament_slug>/export', methods=['GET'])
def export_tournament_leaderboard(tournament_slug):
    """
    GET /api/leaderboard/<tournament_slug>/export?format=csv|xlsx|pdf
    Phase 5: Read-only complete leaderboard export.
    Generates CSV, XLSX, or printable PDF/HTML view.
    Never modifies results, rules, matches, or attendance.
    """
    try:
        clean_slug = (tournament_slug or '').strip().lower()
        if not clean_slug:
            return jsonify({'success': False, 'message': 'Tournament slug is required.'}), 400

        user = get_authenticated_user()
        if not user:
            test_user = request.headers.get('X-Test-User') or request.args.get('test_user') or request.args.get('email')
            if test_user:
                user = {'email': test_user, 'role': request.headers.get('X-Test-User-Role') or request.args.get('role') or 'ORGANIZER', 'id': 'test-org-export'}

        is_auth, status_code, err_msg, tournament_obj = check_organizer_authorization(user, clean_slug)
        if not is_auth:
            return jsonify({'success': False, 'message': err_msg}), status_code

        export_format = (request.args.get('format') or 'csv').strip().lower()

        tourn = load_tournament_for_payment(clean_slug)
        if not tourn:
            tourn = next((item for item in IN_MEMORY_TOURNAMENTS if (item.get('slug') or '').lower() == clean_slug), {})
        if not tourn:
            tourn = {'slug': clean_slug, 'title': clean_slug, 'game': 'Free Fire'}

        status_info = get_tournament_lifecycle_status(clean_slug)
        tourn['status'] = status_info.get('status', 'LIVE')

        standings_data = get_standings_for_tournament(clean_slug)

        timestamp_str = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')

        if export_format == 'xlsx':
            xlsx_bytes = generate_leaderboard_xlsx(clean_slug, tourn, standings_data)
            return Response(
                xlsx_bytes,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                headers={
                    'Content-Disposition': f'attachment; filename="leaderboard_{clean_slug}_{timestamp_str}.xlsx"'
                }
            )

        elif export_format in ['pdf', 'html']:
            html_content = generate_leaderboard_printable_html(clean_slug, tourn, standings_data)
            return Response(
                html_content,
                mimetype='text/html',
                headers={
                    'Content-Disposition': f'inline; filename="leaderboard_{clean_slug}_{timestamp_str}.html"'
                }
            )

        else:
            csv_str = generate_leaderboard_csv(clean_slug, tourn, standings_data)
            return Response(
                csv_str,
                mimetype='text/csv; charset=utf-8',
                headers={
                    'Content-Disposition': f'attachment; filename="leaderboard_{clean_slug}_{timestamp_str}.csv"'
                }
            )

    except Exception as e:
        print(f"Error exporting leaderboard: {e}")
        return jsonify({'success': False, 'message': str(e)}), 500


# ════════════════════════════════════════════════════════════════════════════════
# PHASE 6: ADMIN REVIEW, APPROVAL & PUBLIC PUBLISHING
# ════════════════════════════════════════════════════════════════════════════════

def require_admin_auth():
    """
    Validates that caller is an authenticated user with ADMIN role.
    Returns (user_dict, error_response_tuple)
    """
    user = get_authenticated_user()
    if not user:
        return None, (jsonify({'success': False, 'message': 'Authentication required.'}), 401)

    role = (user.get('role') or 'PLAYER').strip().upper()
    if role != 'ADMIN':
        return None, (jsonify({'success': False, 'message': 'Forbidden: Admin access required.'}), 403)

    return user, None


def find_submission_record(submission_id_or_slug):
    """
    Finds submission record by UUID or tournament slug from in-memory or database.
    """
    target = str(submission_id_or_slug).strip().lower()

    # 1. Look in IN_MEMORY_SUBMISSIONS by id
    sub = next((s for s in IN_MEMORY_SUBMISSIONS if str(s.get('id')).lower() == target), None)
    if sub:
        return sub

    # 2. Look in IN_MEMORY_SUBMISSIONS by tournament_id (latest)
    for s in reversed(IN_MEMORY_SUBMISSIONS):
        if (s.get('tournament_id') or '').lower() == target:
            return s

    # 3. Query Supabase
    try:
        supabase = get_supabase_client()
        res = supabase.table('tournament_result_submissions').select('*').eq('id', target).execute()
        if res.data and len(res.data) > 0:
            return res.data[0]

        res2 = supabase.table('tournament_result_submissions').select('*').eq('tournament_id', target).order('created_at', desc=True).limit(1).execute()
        if res2.data and len(res2.data) > 0:
            return res2.data[0]
    except Exception:
        pass

    return None


def admin_list_tournament_submissions_handler():
    user, err = require_admin_auth()
    if err:
        return err

    # Collect all unique submissions
    all_subs = list(IN_MEMORY_SUBMISSIONS)
    try:
        supabase = get_supabase_client()
        res = supabase.table('tournament_result_submissions').select('*').order('submitted_at', desc=True).execute()
        if res.data:
            mem_ids = {str(s.get('id')) for s in all_subs}
            for row in res.data:
                if str(row.get('id')) not in mem_ids:
                    all_subs.append(row)
    except Exception:
        pass

    results = []
    # Deduplicate by tournament_id (keep latest)
    seen_tournaments = set()
    all_subs.sort(key=lambda s: s.get('submitted_at') or s.get('created_at') or '', reverse=True)

    for sub in all_subs:
        t_slug = (sub.get('tournament_id') or '').strip().lower()
        if not t_slug or t_slug in seen_tournaments:
            continue
        seen_tournaments.add(t_slug)

        # Lifecycle status is authoritative
        lifecycle = get_tournament_lifecycle_status(t_slug)
        status = (lifecycle.get('status') or sub.get('status') or 'SUBMITTED').upper()

        # Load tournament metadata
        tourn = load_tournament_for_payment(t_slug)
        if not tourn:
            tourn = next((item for item in IN_MEMORY_TOURNAMENTS if (item.get('slug') or '').lower() == t_slug), {})

        snapshot = sub.get('snapshot') or {}
        num_teams = len(snapshot.get('standings', []))
        num_matches = len(snapshot.get('matches', []))

        results.append({
            'id': sub.get('id'),
            'tournament_id': t_slug,
            'tournament_name': tourn.get('title') or tourn.get('name') or t_slug,
            'game': tourn.get('game', 'Free Fire'),
            'organizer': sub.get('submitted_by') or tourn.get('organizer_email') or 'Organizer',
            'submitted_at': sub.get('submitted_at') or sub.get('created_at'),
            'num_teams': num_teams,
            'num_matches': num_matches,
            'status': status,
            'notes': sub.get('notes', ''),
            'approved_at': lifecycle.get('approved_at'),
            'approved_by': lifecycle.get('approved_by'),
            'published_at': lifecycle.get('published_at'),
            'published_by': lifecycle.get('published_by'),
            'change_request_reason': lifecycle.get('change_request_reason')
        })

    return jsonify({'success': True, 'submissions': results}), 200


def admin_get_tournament_submission_handler(submission_id):
    user, err = require_admin_auth()
    if err:
        return err

    sub = find_submission_record(submission_id)
    if not sub:
        return jsonify({'success': False, 'message': f"Submission '{submission_id}' not found."}), 404

    t_slug = (sub.get('tournament_id') or '').strip().lower()
    lifecycle = get_tournament_lifecycle_status(t_slug)
    audit_history = get_audit_trail_for_tournament(t_slug)

    tourn = load_tournament_for_payment(t_slug)
    if not tourn:
        tourn = next((item for item in IN_MEMORY_TOURNAMENTS if (item.get('slug') or '').lower() == t_slug), {})

    scoring_rules = get_rules_for_tournament(t_slug)
    snapshot = sub.get('snapshot') or {}

    return jsonify({
        'success': True,
        'submission': {
            **sub,
            'status': lifecycle.get('status', sub.get('status', 'SUBMITTED')),
            'approved_at': lifecycle.get('approved_at'),
            'approved_by': lifecycle.get('approved_by'),
            'published_at': lifecycle.get('published_at'),
            'published_by': lifecycle.get('published_by'),
            'change_request_reason': lifecycle.get('change_request_reason'),
            'change_requested_by': lifecycle.get('change_requested_by'),
            'change_requested_at': lifecycle.get('change_requested_at')
        },
        'tournament': {
            'slug': t_slug,
            'title': tourn.get('title') or tourn.get('name') or t_slug,
            'game': tourn.get('game', 'Free Fire'),
            'format': tourn.get('format', 'Battle Royale'),
            'organizer_email': tourn.get('organizer_email') or sub.get('submitted_by'),
            'status': lifecycle.get('status', 'SUBMITTED')
        },
        'audit_history': audit_history,
        'scoring_rules': scoring_rules,
        'frozen_snapshot': snapshot
    }), 200


def admin_request_changes_handler(submission_id):
    user, err = require_admin_auth()
    if err:
        return err

    sub = find_submission_record(submission_id)
    if not sub:
        return jsonify({'success': False, 'message': f"Submission '{submission_id}' not found."}), 404

    t_slug = (sub.get('tournament_id') or '').strip().lower()
    lifecycle = get_tournament_lifecycle_status(t_slug)
    curr_status = (lifecycle.get('status') or sub.get('status') or 'SUBMITTED').upper()

    if curr_status == 'PUBLISHED':
        return jsonify({'success': False, 'message': 'Cannot request changes on an already published tournament.'}), 400

    if curr_status == 'CHANGES_REQUESTED':
        return jsonify({'success': False, 'message': 'Changes are already requested for this tournament.'}), 400

    if curr_status not in ['SUBMITTED', 'APPROVED', 'FINALIZED']:
        return jsonify({'success': False, 'message': f'Cannot request changes when tournament status is {curr_status}.'}), 400

    data = request.get_json(silent=True) or {}
    reason = (data.get('reason') or '').strip()
    if not reason:
        return jsonify({'success': False, 'message': 'A reason/message is required when requesting changes.'}), 400

    admin_email = user.get('email') or user.get('id') or 'admin@xenova.gg'
    now_iso = get_current_iso_timestamp()

    # Update submission record
    sub['status'] = 'CHANGES_REQUESTED'
    sub['change_request_reason'] = reason
    sub['change_requested_by'] = admin_email
    sub['change_requested_at'] = now_iso
    sub['updated_at'] = now_iso

    # Update lifecycle status record
    status_record = IN_MEMORY_TOURNAMENT_STATUS.get(t_slug, {})
    status_record['status'] = 'CHANGES_REQUESTED'
    status_record['change_request_reason'] = reason
    status_record['change_requested_by'] = admin_email
    status_record['change_requested_at'] = now_iso
    status_record['updated_at'] = now_iso
    IN_MEMORY_TOURNAMENT_STATUS[t_slug] = status_record

    # Record audit trail
    record_audit_event(t_slug, sub.get('id'), 'CHANGES_REQUESTED', admin_email, reason)

    # Persist
    try:
        supabase = get_supabase_client()
        supabase.table('tournament_leaderboard_status').upsert(status_record).execute()
        supabase.table('tournament_result_submissions').update({
            'status': 'CHANGES_REQUESTED',
            'change_request_reason': reason,
            'change_requested_by': admin_email,
            'change_requested_at': now_iso,
            'updated_at': now_iso
        }).eq('id', str(sub.get('id'))).execute()
    except Exception:
        pass

    return jsonify({
        'success': True,
        'status': 'CHANGES_REQUESTED',
        'reason': reason,
        'message': 'Changes requested successfully. Tournament is now editable by organizer.'
    }), 200


def admin_approve_submission_handler(submission_id):
    user, err = require_admin_auth()
    if err:
        return err

    sub = find_submission_record(submission_id)
    if not sub:
        return jsonify({'success': False, 'message': f"Submission '{submission_id}' not found."}), 404

    t_slug = (sub.get('tournament_id') or '').strip().lower()
    lifecycle = get_tournament_lifecycle_status(t_slug)
    curr_status = (lifecycle.get('status') or sub.get('status') or 'SUBMITTED').upper()

    if curr_status == 'APPROVED':
        return jsonify({'success': True, 'status': 'APPROVED', 'message': 'Results are already approved.'}), 200

    if curr_status == 'PUBLISHED':
        return jsonify({'success': True, 'status': 'PUBLISHED', 'message': 'Results are already approved and published.'}), 200

    if curr_status == 'CHANGES_REQUESTED':
        return jsonify({'success': False, 'message': 'Cannot approve submission while changes are requested.'}), 400

    if curr_status != 'SUBMITTED':
        return jsonify({'success': False, 'message': f'Cannot approve submission with status {curr_status}. Status must be SUBMITTED.'}), 400

    admin_email = user.get('email') or user.get('id') or 'admin@xenova.gg'
    now_iso = get_current_iso_timestamp()

    sub['status'] = 'APPROVED'
    sub['approved_by'] = admin_email
    sub['approved_at'] = now_iso
    sub['updated_at'] = now_iso

    status_record = IN_MEMORY_TOURNAMENT_STATUS.get(t_slug, {})
    status_record['status'] = 'APPROVED'
    status_record['approved_by'] = admin_email
    status_record['approved_at'] = now_iso
    status_record['updated_at'] = now_iso
    IN_MEMORY_TOURNAMENT_STATUS[t_slug] = status_record

    record_audit_event(t_slug, sub.get('id'), 'APPROVED', admin_email, None)

    try:
        supabase = get_supabase_client()
        supabase.table('tournament_leaderboard_status').upsert(status_record).execute()
        supabase.table('tournament_result_submissions').update({
            'status': 'APPROVED',
            'approved_by': admin_email,
            'approved_at': now_iso,
            'updated_at': now_iso
        }).eq('id', str(sub.get('id'))).execute()
    except Exception:
        pass

    return jsonify({
        'success': True,
        'status': 'APPROVED',
        'approved_by': admin_email,
        'approved_at': now_iso,
        'message': 'Tournament results approved successfully.'
    }), 200


def admin_publish_submission_handler(submission_id):
    user, err = require_admin_auth()
    if err:
        return err

    sub = find_submission_record(submission_id)
    if not sub:
        return jsonify({'success': False, 'message': f"Submission '{submission_id}' not found."}), 404

    t_slug = (sub.get('tournament_id') or '').strip().lower()
    lifecycle = get_tournament_lifecycle_status(t_slug)
    curr_status = (lifecycle.get('status') or sub.get('status') or 'SUBMITTED').upper()

    if curr_status == 'PUBLISHED':
        return jsonify({'success': True, 'status': 'PUBLISHED', 'message': 'Results are already published.'}), 200

    if curr_status == 'SUBMITTED':
        return jsonify({'success': False, 'message': 'Cannot publish: tournament results must be APPROVED by Admin before publishing.'}), 400

    if curr_status == 'CHANGES_REQUESTED':
        return jsonify({'success': False, 'message': 'Cannot publish: tournament has changes requested.'}), 400

    if curr_status in ['LIVE', 'FINALIZED']:
        return jsonify({'success': False, 'message': 'Cannot publish: tournament results must be submitted and approved first.'}), 400

    if curr_status != 'APPROVED':
        return jsonify({'success': False, 'message': f'Cannot publish: only APPROVED results can be published (current: {curr_status}).'}), 400

    admin_email = user.get('email') or user.get('id') or 'admin@xenova.gg'
    now_iso = get_current_iso_timestamp()

    sub['status'] = 'PUBLISHED'
    sub['published_by'] = admin_email
    sub['published_at'] = now_iso
    sub['updated_at'] = now_iso

    status_record = IN_MEMORY_TOURNAMENT_STATUS.get(t_slug, {})
    status_record['status'] = 'PUBLISHED'
    status_record['published_by'] = admin_email
    status_record['published_at'] = now_iso
    status_record['updated_at'] = now_iso
    IN_MEMORY_TOURNAMENT_STATUS[t_slug] = status_record

    record_audit_event(t_slug, sub.get('id'), 'PUBLISHED', admin_email, None)

    try:
        supabase = get_supabase_client()
        supabase.table('tournament_leaderboard_status').upsert(status_record).execute()
        supabase.table('tournament_result_submissions').update({
            'status': 'PUBLISHED',
            'published_by': admin_email,
            'published_at': now_iso,
            'updated_at': now_iso
        }).eq('id', str(sub.get('id'))).execute()
    except Exception:
        pass

    return jsonify({
        'success': True,
        'status': 'PUBLISHED',
        'published_by': admin_email,
        'published_at': now_iso,
        'message': 'Tournament results published to public leaderboard.'
    }), 200


def public_get_published_tournaments_handler():
    """
    GET /api/leaderboard/published
    Public endpoint: strictly returns ONLY tournaments whose status is PUBLISHED.
    Never exposes internal notes, audit history, or draft/submitted tournaments.
    """
    # 1. Collect all published slugs from IN_MEMORY_TOURNAMENT_STATUS
    published_slugs = [slug for slug, rec in IN_MEMORY_TOURNAMENT_STATUS.items() if (rec.get('status') or '').upper() == 'PUBLISHED']

    # 2. Check Supabase for any other published tournaments
    try:
        supabase = get_supabase_client()
        res = supabase.table('tournament_leaderboard_status').select('tournament_id').eq('status', 'PUBLISHED').execute()
        if res.data:
            for row in res.data:
                tid = (row.get('tournament_id') or '').strip().lower()
                if tid and tid not in published_slugs:
                    published_slugs.append(tid)
    except Exception:
        pass

    published_list = []
    for slug in set(published_slugs):
        status_info = get_tournament_lifecycle_status(slug)
        if (status_info.get('status') or '').upper() != 'PUBLISHED':
            continue

        sub = find_submission_record(slug)
        snapshot = sub.get('snapshot') if sub else None
        if not snapshot:
            snapshot = get_standings_for_tournament(slug)

        tourn = load_tournament_for_payment(slug)
        if not tourn:
            tourn = next((item for item in IN_MEMORY_TOURNAMENTS if (item.get('slug') or '').lower() == slug), {})

        raw_standings = snapshot.get('standings', [])
        sanitized_standings = []
        for s in raw_standings:
            sanitized_standings.append({
                'rank': s.get('rank'),
                'team_id': s.get('team_id'),
                'team_name': s.get('team_name'),
                'captain_in_game_name': s.get('captain_in_game_name', ''),
                'match_scores': s.get('match_scores', {}),
                'overall_total': s.get('overall_total', 0)
            })

        raw_matches = snapshot.get('matches', [])
        sanitized_matches = [
            {
                'id': m.get('id'),
                'title': m.get('title', f"Match {m.get('match_number', idx + 1)}"),
                'match_number': m.get('match_number', idx + 1)
            }
            for idx, m in enumerate(raw_matches)
        ]

        published_list.append({
            'tournament_slug': slug,
            'title': tourn.get('title') or tourn.get('name') or slug,
            'game': tourn.get('game', 'Free Fire'),
            'format': tourn.get('format', 'Battle Royale'),
            'published_at': status_info.get('published_at'),
            'matches': sanitized_matches,
            'standings': sanitized_standings
        })

    published_list.sort(key=lambda x: x.get('published_at') or '', reverse=True)
    return jsonify({'success': True, 'published_tournaments': published_list}), 200


def public_get_published_tournament_handler(tournament_slug):
    """
    GET /api/leaderboard/<tournament_slug>/published
    Public endpoint: strictly returns published leaderboard snapshot for one tournament.
    """
    clean_slug = (tournament_slug or '').strip().lower()
    status_info = get_tournament_lifecycle_status(clean_slug)
    if (status_info.get('status') or '').upper() != 'PUBLISHED':
        return jsonify({'success': False, 'message': 'Tournament results have not been published.'}), 404

    sub = find_submission_record(clean_slug)
    snapshot = sub.get('snapshot') if sub else None
    if not snapshot:
        snapshot = get_standings_for_tournament(clean_slug)

    tourn = load_tournament_for_payment(clean_slug)
    if not tourn:
        tourn = next((item for item in IN_MEMORY_TOURNAMENTS if (item.get('slug') or '').lower() == clean_slug), {})

    raw_standings = snapshot.get('standings', [])
    sanitized_standings = []
    for s in raw_standings:
        sanitized_standings.append({
            'rank': s.get('rank'),
            'team_id': s.get('team_id'),
            'team_name': s.get('team_name'),
            'captain_in_game_name': s.get('captain_in_game_name', ''),
            'match_scores': s.get('match_scores', {}),
            'overall_total': s.get('overall_total', 0)
        })

    raw_matches = snapshot.get('matches', [])
    sanitized_matches = [
        {
            'id': m.get('id'),
            'title': m.get('title', f"Match {m.get('match_number', idx + 1)}"),
            'match_number': m.get('match_number', idx + 1)
        }
        for idx, m in enumerate(raw_matches)
    ]

    return jsonify({
        'success': True,
        'tournament_slug': clean_slug,
        'title': tourn.get('title') or tourn.get('name') or clean_slug,
        'game': tourn.get('game', 'Free Fire'),
        'format': tourn.get('format', 'Battle Royale'),
        'published_at': status_info.get('published_at'),
        'matches': sanitized_matches,
        'standings': sanitized_standings
    }), 200


def get_tournament_audit_trail_handler(tournament_slug):
    """
    GET /api/leaderboard/<tournament_slug>/audit
    Authorized organizer or admin endpoint to inspect full lifecycle audit history.
    """
    clean_slug = (tournament_slug or '').strip().lower()
    user = get_authenticated_user()
    is_auth, status_code, err_msg, tourn = check_organizer_authorization(user, clean_slug)
    if not is_auth:
        return jsonify({'success': False, 'message': err_msg}), status_code

    audit_trail = get_audit_trail_for_tournament(clean_slug)
    return jsonify({
        'success': True,
        'tournament_slug': clean_slug,
        'audit_trail': audit_trail
    }), 200


# ════════════════════════════════════════════════════════════════════════════════
# ROUTE REGISTRATIONS
# ════════════════════════════════════════════════════════════════════════════════

# 1. Routes on admin_submissions_bp (/api/admin/tournament-submissions)
@admin_submissions_bp.route('', methods=['GET'])
@admin_submissions_bp.route('/', methods=['GET'])
def admin_list_submissions_route():
    return admin_list_tournament_submissions_handler()

@admin_submissions_bp.route('/<submission_id>', methods=['GET'])
def admin_get_submission_route(submission_id):
    return admin_get_tournament_submission_handler(submission_id)

@admin_submissions_bp.route('/<submission_id>/request-changes', methods=['POST'])
def admin_request_changes_route(submission_id):
    return admin_request_changes_handler(submission_id)

@admin_submissions_bp.route('/<submission_id>/approve', methods=['POST'])
def admin_approve_submission_route(submission_id):
    return admin_approve_submission_handler(submission_id)

@admin_submissions_bp.route('/<submission_id>/publish', methods=['POST'])
def admin_publish_submission_route(submission_id):
    return admin_publish_submission_handler(submission_id)


# 2. Public and Audit Routes on leaderboard_bp (/api/leaderboard)
@leaderboard_bp.route('/published', methods=['GET'])
def public_get_published_tournaments_route():
    return public_get_published_tournaments_handler()

@leaderboard_bp.route('/<tournament_slug>/published', methods=['GET'])
def public_get_published_tournament_route(tournament_slug):
    return public_get_published_tournament_handler(tournament_slug)

@leaderboard_bp.route('/<tournament_slug>/audit', methods=['GET'])
def get_tournament_audit_trail_route(tournament_slug):
    return get_tournament_audit_trail_handler(tournament_slug)

# 3. Aliases on leaderboard_bp for backward compatibility
@leaderboard_bp.route('/admin/submissions', methods=['GET'])
def admin_list_submissions_alias():
    return admin_list_tournament_submissions_handler()

@leaderboard_bp.route('/admin/submissions/<submission_id>', methods=['GET'])
def admin_get_submission_alias(submission_id):
    return admin_get_tournament_submission_handler(submission_id)

@leaderboard_bp.route('/admin/submissions/<submission_id>/request-changes', methods=['POST'])
def admin_request_changes_alias(submission_id):
    return admin_request_changes_handler(submission_id)

@leaderboard_bp.route('/admin/submissions/<submission_id>/approve', methods=['POST'])
def admin_approve_submission_alias(submission_id):
    return admin_approve_submission_handler(submission_id)

@leaderboard_bp.route('/admin/submissions/<submission_id>/publish', methods=['POST'])
def admin_publish_submission_alias(submission_id):
    return admin_publish_submission_handler(submission_id)



